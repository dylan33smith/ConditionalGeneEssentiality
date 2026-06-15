"""Build the conditional-gene-essentiality lab-meeting deck (python-pptx).

Narrative: Act 1 = the T-regime (predict individual gene fitness on held-out
whole genomes) and why it forced a pivot; Act 2 = the R-regime ranking work and
its findings; Act 3 = synthesis + conclusion. Undergrad speaker notes on every slide.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ----------------------------------------------------------------------------- palette
INK     = "0A2A33"   # deep teal-charcoal (dark backgrounds)
INK2    = "103A45"   # lifted ink (cards on dark)
TEAL    = "0E7C86"   # primary
SEAFOAM = "2A9D8F"   # secondary
MINT    = "8AD9CF"   # light accent
AMBER   = "E0913A"   # sharp accent (negatives / kNN / callouts)
CREAM   = "F4F1EA"   # light background
PAPER   = "FFFFFF"   # cards
INKTX   = "12333B"   # dark text on light
MUTED   = "5E6E73"   # muted text
LINE    = "DBD5C9"   # hairlines on cream

HEAD = "Georgia"
BODY = "Calibri"
EMU_IN = 914400
PW, PH = 13.333, 7.5

prs = Presentation()
prs.slide_width = Emu(int(PW * EMU_IN))
prs.slide_height = Emu(int(PH * EMU_IN))
BLANK = prs.slide_layouts[6]


def rgb(h): return RGBColor.from_string(h)


def slide(bg=CREAM):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = rgb(bg); r.line.fill.background()
    r.shadow.inherit = False
    s.shapes._spTree.remove(r._element); s.shapes._spTree.insert(2, r._element)
    return s


def box(s, x, y, w, h, color=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE,
        radius=None, shadow=False):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if color is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = rgb(color)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line); sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    if shadow:
        el = sp._element.spPr
        ef = el.makeelement(qn('a:effectLst'), {}); el.append(ef)
        sh = ef.makeelement(qn('a:outerShdw'),
                            {'blurRad': '90000', 'dist': '38100', 'dir': '5400000',
                             'rotWithShape': '0'}); ef.append(sh)
        c = sh.makeelement(qn('a:srgbClr'), {'val': '0A2A33'}); sh.append(c)
        a = c.makeelement(qn('a:alpha'), {'val': '16000'}); c.append(a)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0] = radius
        except Exception: pass
    return sp


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space=2, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    if isinstance(runs[0], dict):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space); p.space_before = Pt(0)
        for rn in para:
            r = p.add_run(); r.text = rn["t"]; f = r.font
            f.name = rn.get("font", BODY); f.size = Pt(rn.get("sz", 16))
            f.bold = rn.get("b", False); f.italic = rn.get("i", False)
            f.color.rgb = rgb(rn.get("c", INKTX))
            if "spc" in rn:
                r._r.get_or_add_rPr().set("spc", str(rn["spc"]))
    return tb


def title(s, t, kicker=None, dark=False):
    main = PAPER if dark else INKTX
    box(s, 0.6, 0.62, 0.13, 0.62, color=TEAL if not dark else MINT)
    yy = 0.55
    if kicker:
        text(s, 0.85, yy, 11.8, 0.3, [{"t": kicker.upper(), "sz": 12.5, "b": True,
             "c": SEAFOAM if not dark else MINT, "spc": 220, "font": BODY}])
        yy += 0.34
    text(s, 0.85, yy, 11.9, 0.8, [{"t": t, "sz": 29, "b": True, "c": main, "font": HEAD}])


def footer(s, n, dark=False):
    c = MINT if dark else MUTED
    text(s, 0.6, 7.04, 9, 0.3, [{"t": "Conditional Gene Essentiality  ·  frozen PLM + chemistry → conditional ranking",
         "sz": 9, "c": c, "font": BODY}])
    text(s, 12.0, 7.04, 0.9, 0.3, [{"t": str(n), "sz": 10, "b": True, "c": c, "font": BODY}],
         align=PP_ALIGN.RIGHT)


def chip(s, x, y, label, fill=TEAL, fg=PAPER, w=None, h=0.34, sz=11.5):
    w = w or (0.22 + 0.105 * len(label))
    box(s, x, y, w, h, color=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, x, y + 0.012, w, h - 0.02, [{"t": label, "sz": sz, "b": True, "c": fg, "font": BODY}],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


def bullets(s, x, y, w, items, sz=15, gap=7, lead=TEAL):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(5))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.space_before = Pt(0)
        d = p.add_run(); d.text = "—  "
        d.font.name = BODY; d.font.size = Pt(sz); d.font.bold = True; d.font.color.rgb = rgb(lead)
        if isinstance(it, list):
            for rn in it:
                r = p.add_run(); r.text = rn["t"]; f = r.font
                f.name = BODY; f.size = Pt(rn.get("sz", sz)); f.bold = rn.get("b", False)
                f.italic = rn.get("i", False); f.color.rgb = rgb(rn.get("c", INKTX))
        else:
            r = p.add_run(); r.text = it; f = r.font
            f.name = BODY; f.size = Pt(sz); f.color.rgb = rgb(INKTX)
    return tb


# ============================================================================= 1 TITLE
s = slide(INK)
box(s, 0, 0, PW, 0.16, color=TEAL)
box(s, 0, 7.34, PW, 0.16, color=AMBER)
for dx, dy, r, col in [(11.1, 1.1, 0.16, MINT), (11.9, 1.5, 0.10, SEAFOAM),
        (12.3, 0.9, 0.07, AMBER), (10.7, 1.7, 0.08, SEAFOAM), (12.0, 2.15, 0.13, TEAL)]:
    box(s, dx, dy, r*2, r*2, color=col, shape=MSO_SHAPE.OVAL)
box(s, 0.95, 1.18, 0.13, 0.62, color=MINT)
text(s, 1.2, 1.05, 10, 0.4, [{"t": "INTERNAL LAB MEETING  ·  JUNE 2026", "sz": 13, "b": True,
     "c": MINT, "spc": 260, "font": BODY}])
text(s, 1.2, 1.95, 11.0, 2.2, [
    [{"t": "Can frozen protein-language-model embeddings", "sz": 33, "b": True, "c": PAPER, "font": HEAD}],
    [{"t": "predict conditional gene essentiality?", "sz": 33, "b": True, "c": MINT, "font": HEAD}],
], space=4)
text(s, 1.2, 3.95, 10.8, 1.0, [{"t": "Two attempts on Tn-seq fitness data. First we tried to predict each "
     "gene's fitness on never-seen genomes — it failed instructively. We reframed to ranking a known gene's "
     "stressors, and found a simple chemistry lookup beats every learned model we built.",
     "sz": 15.5, "c": "C9DBDF", "i": True, "font": BODY}])
text(s, 1.2, 5.55, 11, 0.4, [
    [{"t": "T-regime", "sz": 12.5, "b": True, "c": MINT},
     {"t": ": cross-organism fitness regression    →    ", "sz": 12.5, "c": "9FB6BB"},
     {"t": "R-regime", "sz": 12.5, "b": True, "c": MINT},
     {"t": ": within-organism conditional ranking", "sz": 12.5, "c": "9FB6BB"}]])

# ============================================================================= 2 T-GOAL  (NEW)
s = slide(CREAM)
title(s, "Attempt 1 — predict a gene's fitness on a new genome", kicker="The T-regime · cross-organism regression")
bullets(s, 0.85, 1.7, 6.3, [
    [{"t": "The original goal: ", "b": True}, {"t": "given a gene (its ProteomeLM embedding) and a condition "
     "(its chemistry), predict the "}, {"t": "continuous fitness value", "b": True, "c": TEAL},
     {"t": " — a regression task."}],
    [{"t": "The ambitious test — ", "b": True}, {"t": "whole-genome (cross-organism) holdout: "},
     {"t": "train on some bacterial species, test on ", "b": False},
     {"t": "entirely different ones", "b": True, "c": AMBER}, {"t": "."}],
    [{"t": "Why it's hard: ", "b": True}, {"t": "the test genes belong to organisms never seen in training, "
     "so the model must generalize from "}, {"t": "protein sequence alone", "b": True},
     {"t": " — no prior measurements exist for them."}],
    [{"t": "Scored by ", "b": False}, {"t": "RMSE + MAE", "b": True, "c": SEAFOAM},
     {"t": " (average prediction error). The dream: essentiality for a microbe you've never assayed."}],
], sz=14.5, gap=11)
# right schematic card: whole-genome holdout
cx, cy, cw = 7.7, 1.72, 5.0
box(s, cx, cy, cw, 4.4, color=PAPER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, shadow=True)
text(s, cx + 0.35, cy + 0.25, cw - 0.7, 0.35, [{"t": "WHOLE-GENOME HOLDOUT", "sz": 12, "b": True,
     "c": SEAFOAM, "spc": 160}])
text(s, cx + 0.35, cy + 0.62, cw - 0.7, 0.3, [{"t": "TRAIN on many genomes", "sz": 12.5, "b": True, "c": INKTX}])
orgs = ["Keio", "DvH", "Caulo", "Putida", "MR1", "Btheta"]
gx, gy = cx + 0.35, cy + 1.0
for i, o in enumerate(orgs):
    chip(s, gx + (i % 3) * 1.45, gy + (i // 3) * 0.5, o, fill=TEAL, w=1.32, sz=11)
# arrow down
box(s, cx + cw/2 - 0.05, cy + 2.05, 0.1, 0.42, color=MUTED)
ar = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(cx + cw/2 - 0.22),
                        Inches(cy + 2.42), Inches(0.44), Inches(0.26))
ar.rotation = 180; ar.fill.solid(); ar.fill.fore_color.rgb = rgb(MUTED); ar.line.fill.background()
ar.shadow.inherit = False
text(s, cx + 0.35, cy + 2.85, cw - 0.7, 0.3, [{"t": "TEST on a held-out genome", "sz": 12.5, "b": True, "c": INKTX}])
chip(s, cx + 0.35, cy + 3.2, "new species  ·  genes never seen", fill=AMBER, fg=INK, w=cw - 0.7, h=0.42, sz=12.5)
text(s, cx + 0.35, cy + 3.78, cw - 0.7, 0.5, [{"t": "no training history → predict from sequence alone",
     "sz": 11, "i": True, "c": MUTED}])
footer(s, 2)

# ============================================================================= 3 T-PIPELINE (NEW)
s = slide(CREAM)
title(s, "The T pipeline — a pre-registered tier sweep", kicker="Systematic model search")
text(s, 0.85, 1.55, 11.8, 0.55, [
    [{"t": "We didn't train one model and stop. ", "sz": 13.5, "c": INKTX},
     {"t": "Each ‘tier’ isolated ONE design choice, with success criteria fixed in advance and a "
      "decision-ledger entry recording the winner.", "sz": 13.5, "c": INKTX}]])
# governance strip
box(s, 0.85, 2.2, 11.62, 0.62, color="E7EEEC", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
text(s, 1.1, 2.2, 11.2, 0.62, [
    [{"t": "S0–S5  governance & setup:  ", "sz": 12.5, "b": True, "c": TEAL},
     {"t": "reproducibility · data characterization · evaluation trust · split lock · feature contract · training recipe",
      "sz": 12, "c": INKTX}]], anchor=MSO_ANCHOR.MIDDLE)
tiers = [
    ("T1", "Representation", "425-d multihot chemistry wins; fingerprints & metadata don't help", TEAL),
    ("T2", "Fusion", "early-concat shallow MLP beats two-tower & FiLM", SEAFOAM),
    ("T3", "Capacity", "2-layer residual MLP, width 512 — bigger doesn't help", AMBER),
    ("T4", "Optimization", "MSE on raw targets; short constant-LR training", TEAL),
    ("T5", "Embedding", "trainable adapter on ProteomeLM-L8 — first real gain; beats ESM-C", SEAFOAM),
    ("T6", "Chemistry", "multihot still beats Morgan / RDKit / MACCS fingerprints", AMBER),
]
cw2, ch2, gapx, gapy = 3.68, 1.78, 0.29, 0.2
x0, y0 = 0.85, 3.05
for i, (code, name, desc, col) in enumerate(tiers):
    cx = x0 + (i % 3) * (cw2 + gapx)
    cy = y0 + (i // 3) * (ch2 + gapy)
    box(s, cx, cy, cw2, ch2, color=PAPER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, shadow=True)
    box(s, cx, cy, 0.13, ch2, color=col)
    text(s, cx + 0.3, cy + 0.2, 1.2, 0.5, [{"t": code, "sz": 21, "b": True, "c": col, "font": HEAD}])
    text(s, cx + 1.15, cy + 0.28, cw2 - 1.3, 0.4, [{"t": name, "sz": 14.5, "b": True, "c": INKTX}])
    text(s, cx + 0.3, cy + 0.86, cw2 - 0.55, 0.85, [{"t": desc, "sz": 11.5, "c": MUTED}])
footer(s, 3)

# ============================================================================= 4 T-VERDICT / PIVOT (was cross-org)
s = slide(CREAM)
title(s, "The verdict: optimized error, but no real signal", kicker="The turning point — why we reframed")
bullets(s, 0.85, 1.72, 6.7, [
    [{"t": "RMSE/MAE were optimized fine across T1–T6 — but the ", "c": INKTX},
     {"t": "meaningful quantity failed.", "b": True, "c": AMBER}],
    [{"t": "Within-gene ranking", "b": True, "c": TEAL}, {"t": " of conditions: "},
     {"t": "≈ 0.045 Spearman", "b": True, "c": AMBER}, {"t": " vs a ~0.43 replicate noise floor — "
     "statistically indistinguishable from random.", "c": INKTX}],
    [{"t": "Why RMSE misled us: it is ", "c": INKTX}, {"t": "gene-mean-dominated", "b": True, "c": INKTX},
     {"t": " — it rewards predicting each gene's baseline level, not the gene×condition interaction."}],
    [{"t": "Frozen embeddings carry ", "c": INKTX}, {"t": "no transferable conditional signal", "b": True, "c": INKTX},
     {"t": "; ProteomeLM beat raw ESM-C by only RMSE ~0.006."}],
], sz=14, gap=10)
box(s, 8.0, 1.78, 4.7, 3.55, color=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, shadow=True)
text(s, 8.35, 2.1, 4.0, 0.4, [{"t": "THE HONEST FRAME", "sz": 11.5, "b": True, "c": MINT, "spc": 180}])
text(s, 8.35, 2.6, 4.05, 2.6, [
    [{"t": "Predicting the level", "sz": 15, "b": True, "c": PAPER}],
    [{"t": "of fitness is easy and gene-mean-dominated.", "sz": 12.5, "c": "C4D6D9"}],
    [{"t": " ", "sz": 5}],
    [{"t": "Predicting the conditional ordering", "sz": 15, "b": True, "c": MINT}],
    [{"t": "is the real, hard problem — and cross-organism transfer of it is ≈ 0.", "sz": 12.5, "c": "C4D6D9"}],
], space=4)
box(s, 0.85, 5.7, 11.65, 0.66, color=AMBER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
text(s, 1.1, 5.7, 11.2, 0.66, [
    [{"t": "→ The pivot:  ", "sz": 14.5, "b": True, "c": INK},
     {"t": "narrow the claim to WITHIN-organism, and switch from predicting values to ranking conditions. "
      "That defines the rest of the talk.", "sz": 14, "c": INK}]], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 4)

# ============================================================================= 5 R-QUESTION
s = slide(CREAM)
title(s, "The reframed question, stated precisely", kicker="The R-regime · motivation")
text(s, 0.85, 1.65, 7.0, 2.6, [
    [{"t": "Within a known organism", "sz": 17, "b": True, "c": TEAL},
     {"t": ", can we rank a gene's conditional", "sz": 17, "c": INKTX}],
    [{"t": "essentiality across ", "sz": 17, "c": INKTX},
     {"t": "entirely novel conditions", "sz": 17, "b": True, "c": TEAL},
     {"t": " — better", "sz": 17, "c": INKTX}],
    [{"t": "than a chemistry-similarity baseline that uses condition", "sz": 17, "c": INKTX}],
    [{"t": "features but learns no gene-specific interaction?", "sz": 17, "c": INKTX}],
], space=4)
text(s, 0.85, 3.55, 7.0, 0.9, [{"t": "Framed for the biologist: “find the top stressors for this gene.”",
     "sz": 14.5, "i": True, "c": MUTED}])
box(s, 8.3, 1.7, 4.4, 4.5, color=PAPER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, shadow=True)
text(s, 8.65, 1.95, 3.8, 0.4, [{"t": "THE INGREDIENTS", "sz": 11.5, "b": True, "c": SEAFOAM, "spc": 200}])
rows = [("Gene", "frozen ProteomeLM-L8 embedding (1152-d) — sequence/proteome context", TEAL),
        ("Condition", "media + stressor chemistry (425-d multihot over a canonical vocab)", SEAFOAM),
        ("Label", "Tn-seq fitness; relevance = max(0, −fit) = “this stresses the gene”", AMBER),
        ("Target", "within-gene ordering of conditions (Spearman + NDCG@5)", INK)]
yy = 2.45
for name, desc, col in rows:
    box(s, 8.65, yy + 0.04, 0.12, 0.78, color=col)
    text(s, 8.95, yy, 3.55, 0.35, [{"t": name, "sz": 14, "b": True, "c": INKTX}])
    text(s, 8.95, yy + 0.32, 3.55, 0.6, [{"t": desc, "sz": 11.5, "c": MUTED}])
    yy += 0.92
footer(s, 5)

# ============================================================================= 6 R-SETUP
s = slide(CREAM)
title(s, "Setup: data and the cold-start split", kicker="Task design")
cards = [
    ("27.4M", "gene × experiment fitness rows", "Fitness Browser / RB-TnSeq (feba.db)", TEAL),
    ("23", "organisms with replicate structure", "for the full within-org evaluation", SEAFOAM),
    ("0 / 78", "train↔val condition overlap", "held-out conditions are 100% disjoint", AMBER),
]
x = 0.85
for big, lab, sub, col in cards:
    box(s, x, 1.7, 3.78, 1.7, color=PAPER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07, shadow=True)
    box(s, x, 1.7, 0.13, 1.7, color=col)
    text(s, x + 0.34, 1.86, 3.3, 0.8, [{"t": big, "sz": 33, "b": True, "c": col, "font": HEAD}])
    text(s, x + 0.34, 2.62, 3.3, 0.4, [{"t": lab, "sz": 13, "b": True, "c": INKTX}])
    text(s, x + 0.34, 2.97, 3.3, 0.4, [{"t": sub, "sz": 10.5, "c": MUTED}])
    x += 3.97
text(s, 0.85, 3.95, 11.8, 0.4, [{"t": "Why the split makes this hard", "sz": 16, "b": True, "c": TEAL, "font": HEAD}])
bullets(s, 0.85, 4.45, 11.7, [
    [{"t": "Condition-holdout: ", "b": True}, {"t": "we hold out whole conditions, replicate-grouped and "
     "stratified by stressor group — val conditions are never seen in train."}],
    [{"t": "Cold columns: ", "b": True}, {"t": "a held-out condition has zero observed entries, so vanilla "
     "collaborative matrix factorization cannot place it. You genuinely need condition side-features (chemistry)."}],
    [{"t": "So this is ", "b": False}, {"t": "inductive (cold-start) matrix completion with side information", "b": True,
     "c": TEAL}, {"t": " — harder and more defensible than warm-cell completion."}],
], sz=14, gap=8)
footer(s, 6)

# ============================================================================= 7 R-EVAL
s = slide(CREAM)
title(s, "How we measure — and the baseline ladder", kicker="Evaluation contract")
text(s, 0.85, 1.6, 5.9, 0.4, [{"t": "Co-primary metrics", "sz": 15, "b": True, "c": TEAL, "font": HEAD}])
bullets(s, 0.85, 2.05, 6.0, [
    [{"t": "within-gene Spearman", "b": True}, {"t": "  +  "}, {"t": "NDCG@5", "b": True},
     {"t": "  (“top stressors”)"}],
    [{"t": "Hierarchical org→gene bootstrap; promotion needs "}, {"t": "ΔNDCG@5 ≳ 0.026", "b": True, "c": AMBER},
     {"t": " over the gate, CIs disjoint."}],
], sz=13.5, gap=7)
text(s, 0.85, 3.25, 6.0, 0.4, [{"t": "Four reference bars", "sz": 15, "b": True, "c": TEAL, "font": HEAD}])
bullets(s, 0.85, 3.7, 6.0, [
    [{"t": "chem-NULL", "b": True, "c": MUTED}, {"t": " — population profile, no gene specificity"}],
    [{"t": "linear-MF", "b": True, "c": SEAFOAM}, {"t": " — a learned, fitness-aware gene latent"}],
    [{"t": "chem-kNN", "b": True, "c": AMBER}, {"t": " — the gene's OWN history at near chemistries (the gate)"}],
    [{"t": "ceiling", "b": True, "c": INK}, {"t": " — biological-replicate agreement"}],
], sz=13, gap=5)
rows = [["method", "Spearman", "NDCG@5"],
        ["chem-NULL", "0.02", "0.31"],
        ["deep model (frozen emb + chem)", "0.13", "0.42"],
        ["linear-MF (learned latent)", "0.14", "0.43"],
        ["chem-kNN  ◂ the baseline to beat", "0.24", "0.485"],
        ["replicate ceiling", "0.39", "0.66"]]
tx, ty, tw, rh = 7.15, 1.75, 5.5, 0.62
for i, row in enumerate(rows):
    yy = ty + i * rh
    if i == 0:
        bg, fg, bold = INK, PAPER, True
    elif "chem-kNN" in row[0]:
        bg, fg, bold = AMBER, INK, True
    elif "ceiling" in row[0]:
        bg, fg, bold = "E7EEEC", INKTX, True
    else:
        bg, fg, bold = PAPER, INKTX, False
    box(s, tx, yy, tw, rh - 0.07, color=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(s, tx + 0.22, yy, 3.5, rh - 0.07, [{"t": row[0], "sz": 12.5, "b": bold, "c": fg}],
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, tx + 3.55, yy, 0.95, rh - 0.07, [{"t": row[1], "sz": 12.5, "b": bold, "c": fg}],
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    text(s, tx + 4.5, yy, 0.95, rh - 0.07, [{"t": row[2], "sz": 12.5, "b": bold, "c": fg}],
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
footer(s, 7)

# ============================================================================= 8 FINDING 1
s = slide(CREAM)
title(s, "Finding 1 — encoder, objective, capacity: all fail", kicker="Three independent levers")
cards = [
    ("ENCODER", "R1", "multihot vs Morgan / RDKit /\nMACCS fingerprints", "6 arms × 3 seeds × 23 orgs",
     "best 0.422 — fingerprints don't help", TEAL),
    ("OBJECTIVE", "R-LOSS", "MSE, Huber, RankNet,\nLambdaRank, ListMLE, ApproxNDCG", "6 loss families",
     "best (Huber) 0.435 — ranking losses lose", SEAFOAM),
    ("CAPACITY", "—", "linear-MF (~13k params) vs\ndeep MLP (2.5M params)", "bilinear vs nonlinear",
     "linear ≈ deep — capacity buys nothing", AMBER),
]
x = 0.85
for tag, code, what, scale, verdict, col in cards:
    box(s, x, 1.7, 3.82, 4.0, color=PAPER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, shadow=True)
    box(s, x, 1.7, 3.82, 0.62, color=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    box(s, x, 2.05, 3.82, 0.27, color=col)
    text(s, x + 0.3, 1.7, 2.6, 0.62, [{"t": tag, "sz": 14, "b": True, "c": PAPER, "spc": 120}],
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 2.9, 1.7, 0.7, 0.62, [{"t": code, "sz": 13, "b": True, "c": "FFFFFF", "i": True}],
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
    text(s, x + 0.3, 2.62, 3.25, 1.0, [{"t": what, "sz": 13, "b": True, "c": INKTX}])
    text(s, x + 0.3, 3.75, 3.25, 0.4, [{"t": scale, "sz": 11, "i": True, "c": MUTED}])
    box(s, x + 0.3, 4.25, 3.22, 0.02, color=LINE)
    text(s, x + 0.3, 4.45, 3.25, 1.0, [
        [{"t": "✗  ", "sz": 13, "b": True, "c": AMBER}, {"t": verdict, "sz": 12.5, "b": True, "c": INKTX}]])
    x += 3.97
text(s, 0.85, 5.95, 11.8, 0.7, [
    [{"t": "Every global parametric model plateaus at NDCG@5 ≈ 0.42–0.44.  ", "sz": 14.5, "b": True, "c": INKTX},
     {"t": "None beats the chem-kNN gate (0.485).", "sz": 14.5, "b": True, "c": AMBER}]])
footer(s, 8)

# ============================================================================= 9 MECHANISM
s = slide(INK)
title(s, "The mechanism: local vs global", kicker="Why the lookup wins", dark=True)
bullets(s, 0.85, 1.75, 6.5, [
    [{"t": "The population / cross-gene pattern carries ", "c": "DCE8EA"},
     {"t": "≈ 0 within-gene signal", "b": True, "c": MINT}, {"t": " (chem-NULL).", "c": "DCE8EA"}],
    [{"t": "So all the signal lives in each gene's ", "c": "DCE8EA"},
     {"t": "own deviations", "b": True, "c": MINT}, {"t": " from the bulk.", "c": "DCE8EA"}],
    [{"t": "A global model minimizing average error over ~54k genes captures the bulk and ", "c": "DCE8EA"},
     {"t": "averages those idiosyncratic deviations away", "b": True, "c": AMBER}, {"t": ".", "c": "DCE8EA"}],
    [{"t": "kNN ", "b": True, "c": MINT}, {"t": "preserves", "c": "DCE8EA"},
     {"t": " them — it reads each gene's actual history and interpolates by chemistry similarity.", "c": "DCE8EA"}],
], sz=15, gap=11, lead=MINT)
box(s, 7.7, 1.75, 5.0, 4.35, color=INK2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
box(s, 7.7, 1.75, 0.13, 4.35, color=AMBER)
text(s, 8.05, 2.05, 4.4, 0.4, [{"t": "THE ANALOGY", "sz": 11.5, "b": True, "c": AMBER, "spc": 200}])
text(s, 8.05, 2.5, 4.45, 1.4, [
    [{"t": "Global model", "sz": 14.5, "b": True, "c": PAPER}],
    [{"t": "= a doctor predicting a never-seen patient from their “type.”", "sz": 13, "c": "C4D6D9"}],
], space=3)
text(s, 8.05, 4.0, 4.45, 1.8, [
    [{"t": "chem-kNN", "sz": 14.5, "b": True, "c": MINT}],
    [{"t": "= a doctor holding that patient's full chart, predicting from the most similar past situations "
     "for that patient.", "sz": 13, "c": "C4D6D9"}],
], space=3)
text(s, 8.05, 5.55, 4.45, 0.5, [{"t": "When the chart exists and is idiosyncratic, the second wins.",
     "sz": 12, "i": True, "c": MINT}])
footer(s, 9, dark=True)

# ============================================================================= 10 FINDING 2 HYBRID
s = slide(CREAM)
title(s, "Finding 2 — hybrids don't beat the lookup either", kicker="Global + local fusion")
text(s, 0.85, 1.62, 11.7, 0.7, [
    [{"t": "If the model carries ", "sz": 14.5, "c": INKTX},
     {"t": "complementary", "sz": 14.5, "b": True, "c": TEAL},
     {"t": " signal, a learned hybrid of (global model + chem-kNN) should beat the gate. "
      "We tried four fusions across two rounds.", "sz": 14.5, "c": INKTX}]])
rows = [["fusion", "what it is", "Δ NDCG@5 vs gate", "verdict"],
        ["R-HYBRID-A · static ensemble", "z-scored convex blend, held-out α", "+0.008", "below 0.026 delta"],
        ["R-HYBRID-B · residual", "model predicts kNN's LOO residual", "+0.0002", "tie; sign-flips by seed"],
        ["R-HYBRID-B · retrieval-augmented", "feed gene's k-nearest fits to model", "−0.0059", "regresses"],
        ["R-HYBRID-B · learned gating", "per-cell α(gene, condition)", "−0.0081", "regresses"]]
tx, ty, tw = 0.85, 2.5, 11.65
colw = [3.5, 4.0, 2.1, 2.05]
rh = 0.66
for i, row in enumerate(rows):
    yy = ty + i * rh
    head = i == 0
    bg = INK if head else PAPER
    fg = PAPER if head else INKTX
    box(s, tx, yy, tw, rh - 0.08, color=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, shadow=not head)
    cx = tx + 0.25
    for j, cell in enumerate(row):
        cc = fg; bold = head
        if not head and j == 2:
            cc = TEAL if cell == "+0.008" else AMBER
            bold = True
        if not head and j == 0:
            bold = True
        al = PP_ALIGN.LEFT if j < 2 else PP_ALIGN.CENTER
        text(s, cx if j < 2 else cx - 0.1, yy, colw[j] - 0.1, rh - 0.08,
             [{"t": cell, "sz": 12 if j != 1 else 11.5, "b": bold, "c": cc}],
             anchor=MSO_ANCHOR.MIDDLE, align=al)
        cx += colw[j]
text(s, 0.85, 5.95, 11.8, 0.7, [
    [{"t": "Three orthogonal fusions converge on ≈ 0.  ", "sz": 14.5, "b": True, "c": AMBER},
     {"t": "The complementary signal is real but tiny — not promotable.", "sz": 14.5, "b": True, "c": INKTX}]])
footer(s, 10)

# ============================================================================= 11 FINDING 3 R-CONF
s = slide(CREAM)
title(s, "Finding 3 — and it's NOT just label noise", kicker="R-CONF · confidence stratification")
text(s, 0.85, 1.58, 11.8, 0.55, [
    [{"t": "Objection: “your labels are noisy, the metric is measuring noise.”  We stratified val genes by "
      "measurement confidence (|t|, the Wetmore moderated-t).", "sz": 13.5, "c": INKTX}]])
fw = 8.5; fh = fw / 2.667
s.shapes.add_picture("research_log/figures/r_conf/01_confidence_stratified.png",
                     Inches(0.85), Inches(2.25), width=Inches(fw), height=Inches(fh))
box(s, 0.85, 2.25, fw, fh, line=LINE, lw=1.0)
tx = 9.7
text(s, tx, 2.2, 3.0, 0.4, [{"t": "WHAT IT SHOWS", "sz": 11.5, "b": True, "c": SEAFOAM, "spc": 180}])
bullets(s, tx, 2.6, 3.05, [
    [{"t": "kNN beats the model in ", "c": INKTX}, {"t": "every", "b": True, "c": AMBER},
     {"t": " confidence stratum.", "c": INKTX}],
    [{"t": "Gap narrows (0.055→0.040) but ", "c": INKTX}, {"t": "never closes", "b": True, "c": AMBER},
     {"t": ".", "c": INKTX}],
    [{"t": "Noise is real: ceiling rises 0.59→0.83 with |t|.", "c": INKTX}],
    [{"t": "t-weighted training ", "c": INKTX}, {"t": "Δ ≈ −0.004", "b": True, "c": MUTED},
     {"t": " (no help).", "c": INKTX}],
], sz=12, gap=8)
box(s, 0.85, 6.35, 11.65, 0.62, color="E7EEEC", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
text(s, 1.1, 6.35, 11.2, 0.62, [
    [{"t": "The negative is structural, not a measurement artifact", "sz": 14, "b": True, "c": TEAL},
     {"t": " — the model loses to lookup even on the cleanest-measured genes.", "sz": 14, "c": INKTX}]],
    anchor=MSO_ANCHOR.MIDDLE)
footer(s, 11)

# ============================================================================= 12 DURABLE
s = slide(CREAM)
title(s, "What is durably established", kicker="Summary of findings")
items = [
    ("Modest, real signal", "Conditional essentiality has a within-gene signal, but the ceiling is low — replicates agree only at NDCG@5 ~0.66. A noisy target."),
    ("Local & gene-idiosyncratic", "Population/cross-gene structure ≈ 0 within-gene signal; each gene's own history carries it."),
    ("No cross-org transfer", "Frozen embeddings carry no transferable conditional-response information (Spearman ≈ noise)."),
    ("kNN is the strong baseline", "A chemistry-similarity lookup (0.485) beats every learned global model — across encoders, objectives, capacities, frozen & learned-fitness-aware reps, alone & hybridized."),
    ("Memorization-dominated", "The useful product — rank a known gene's stressors in a known organism — is achievable via lookup. Deep learning, as configured, adds nothing over it."),
    ("Noise-robust negative", "The result holds at every measurement-confidence level (R-CONF); cleaner labels would not change it."),
]
x, y = 0.85, 1.7
cw, ch = 5.78, 1.62
for i, (h, d) in enumerate(items):
    col = [TEAL, SEAFOAM, AMBER, TEAL, SEAFOAM, AMBER][i]
    cx = x + (i % 2) * (cw + 0.25)
    cy = y + (i // 2) * (ch + 0.12)
    box(s, cx, cy, cw, ch, color=PAPER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, shadow=True)
    box(s, cx, cy, 0.13, ch, color=col)
    box(s, cx + 0.32, cy + 0.26, 0.42, 0.42, color=col, shape=MSO_SHAPE.OVAL)
    text(s, cx + 0.32, cy + 0.26, 0.42, 0.42, [{"t": str(i+1), "sz": 15, "b": True, "c": PAPER, "font": HEAD}],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, cx + 0.92, cy + 0.22, cw - 1.1, 0.4, [{"t": h, "sz": 14, "b": True, "c": INKTX}])
    text(s, cx + 0.92, cy + 0.62, cw - 1.15, 0.95, [{"t": d, "sz": 10.8, "c": MUTED}])
footer(s, 12)

# ============================================================================= 13 CONCLUSION
s = slide(INK)
box(s, 0, 0, PW, 0.16, color=TEAL)
box(s, 0, 7.34, PW, 0.16, color=AMBER)
title(s, "The honest scope — and what's left", kicker="Conclusion", dark=True)
text(s, 0.85, 1.75, 11.7, 0.9, [
    [{"t": "The modeling thread is closed with a clean, pre-registered negative. ", "sz": 16, "b": True, "c": PAPER},
     {"t": "That result — a rigorous benchmark showing a simple lookup beats learned models — ", "sz": 16, "c": "DCE8EA"},
     {"t": "is the contribution.", "sz": 16, "b": True, "c": MINT}]])
paths = [
    ("ADOPTED", "Characterization paper", "Publish the rigorous negative: memorization-dominated task, "
     "strong-baseline benchmark, noise-robust. Honest, pre-registered, credible.", MINT, AMBER),
    ("OPTIONAL", "Cold-gene diagnostic", "Quantify how far kNN degrades on unseen genes — the one regime a "
     "global model could help. A figure, not a gate.", "C4D6D9", SEAFOAM),
    ("DEFERRED", "R-EMB moonshot", "Fitness-aware inductive embedding. linear-MF says it won't help warm; "
     "its only target (cold transfer) has a discouraging prior.", "9FB6BB", "6E8186"),
]
x = 0.85
for tag, h, d, tc, badge in paths:
    box(s, x, 2.95, 3.82, 3.1, color=INK2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    box(s, x, 2.95, 3.82, 0.13, color=badge)
    chip(s, x + 0.3, 3.28, tag, fill=badge, fg=INK if tag == "ADOPTED" else PAPER)
    text(s, x + 0.3, 3.85, 3.3, 0.5, [{"t": h, "sz": 16.5, "b": True, "c": PAPER, "font": HEAD}])
    text(s, x + 0.3, 4.45, 3.32, 1.5, [{"t": d, "sz": 12, "c": tc}])
    x += 3.97
footer(s, 13, dark=True)

# ============================================================================= NOTES
NOTES = {
1: [
 "PURPOSE. We tested whether a modern AI model can predict which growth conditions a bacterial gene is most sensitive to. The honest answer is 'a simple lookup does it better.' These notes define the jargon as we go.",
 "Roadmap (two acts): first we tried to predict each gene's exact fitness on never-seen genomes — the 'T-regime'; it failed in an instructive way. So we reframed to RANKING a known gene's stressors — the 'R-regime'. The slides follow that order.",
 "Gene essentiality: a gene is 'essential' if the cell can't grow without it. CONDITIONAL essentiality means a gene is needed only in some situations (a certain nutrient missing, a toxin present), not always — that's what we predict.",
 "Tn-seq (transposon sequencing): a wet-lab method that disables every gene across a population, grows them in some condition, and sequences to see which knockouts became rarer. If disabling a gene hurts growth, that gene mattered there. The readout is a 'fitness' score per gene per condition.",
 "Protein language model (e.g., ProteomeLM): an AI trained on millions of protein sequences, like a language model for proteins. It turns a protein into an 'embedding' — a list of numbers capturing its properties. 'Frozen' = used as-is, never retrained.",
],
2: [
 "This is the project's FIRST attempt, chronologically before the ranking work.",
 "Regression vs ranking: here we tried to predict the actual NUMBER (continuous fitness) — that's 'regression'. Later we switch to predicting just the ORDER of conditions — 'ranking'.",
 "Whole-genome / cross-organism holdout: we trained on some bacterial species and tested on ENTIRELY DIFFERENT species, holding out whole genomes. This is the most ambitious kind of generalization: the test genes belong to organisms never seen in training, so the model has no prior measurements for them and must work from protein sequence alone.",
 "RMSE (root-mean-square error) and MAE (mean absolute error): two ways to measure how far predictions land from the true values; lower is better. 'Co-primary' means we required both to improve.",
 "Why it's the dream: if it worked, you could predict gene essentiality for a brand-new microbe you've never assayed in the lab.",
 "The right-hand schematic: teal chips = species used for training; the amber chip = a held-out species whose genes have no training history.",
],
3: [
 "We didn't train one network and stop. We ran a disciplined, PRE-REGISTERED sequence of experiments ('tiers'), each isolating ONE design choice, with success criteria written down in advance and a 'decision-ledger' entry recording the outcome. That discipline is what makes the eventual negative result trustworthy rather than a case of 'not trying hard enough'.",
 "S0–S5 are setup stages: reproducibility (can we re-run exactly?), data characterization (what's actually in the data?), evaluation trustworthiness (are our metrics and baselines sound?), split lock (freeze the train/test division so we can't cheat), feature contract (freeze how inputs are encoded), and training recipe (fix data-cleaning and which organisms to pool).",
 "T1 Representation: how to encode the condition's chemistry. A simple 'multihot' vector (which chemicals are present) won.",
 "T2 Fusion: how to combine the gene vector and the chemistry vector. Concatenating them early and feeding a small neural net (MLP = multi-layer perceptron) beat fancier designs (two-tower, FiLM).",
 "T3 Capacity: how big/deep the network is. A 2-layer residual MLP of width 512 was enough; bigger didn't help. (A 'residual' connection lets a layer learn an adjustment on top of its input, which helps training.)",
 "T4 Optimization: the loss and training schedule — plain MSE on raw (un-normalized) targets, trained briefly.",
 "T5 Embedding: we questioned the frozen embedding and added a small trainable 'adapter' on top of ProteomeLM's layer 8 — the first real improvement, and ProteomeLM beat another protein model, ESM-C.",
 "T6 Chemistry: we retried molecular 'fingerprints' (structure-based encodings) — multihot still won. Bottom line: we exhaustively optimized the model.",
],
4: [
 "This slide is the turning point. Despite optimizing RMSE/MAE across the whole T pipeline, the biologically meaningful quantity failed.",
 "Within-gene ranking: instead of error on raw values, ask 'did we get the ORDER of conditions right for each gene?' We measure it with Spearman correlation (−1 to +1; +1 = perfect order). We got ≈ 0.045 — essentially random.",
 "Noise floor (a.k.a. ceiling): the score you'd get just from one biological replicate predicting another — the best achievable given measurement noise. Here it's ~0.43, so 0.045 is far below it: a real failure, not a hard-but-okay result.",
 "Why RMSE looked fine yet ranking failed: RMSE is 'gene-mean-dominated' — most of its score comes from getting each gene's overall AVERAGE fitness level right (easy), which says nothing about the gene×condition INTERACTION (the specific, hard signal we actually want).",
 "ESM-C: a different, well-known protein language model we compared against. ProteomeLM beat it by a negligible margin (~0.006 RMSE), reinforcing that sequence embeddings simply don't encode how a gene responds to stress — they're 'fitness-blind'.",
 "The decision: cross-organism (whole-genome) transfer of the conditional signal is ≈ 0. So we made two changes — narrowed the claim from cross-organism to WITHIN-organism, and switched from predicting values (regression) to ranking conditions. That pivot ('R-regime') defines the rest of the talk.",
],
5: [
 "RANK, not predict exact values: a biologist mainly wants the ORDER — which conditions stress this gene most — not the precise number. Ranking is both easier and more useful, so all our R-regime metrics are ranking metrics.",
 "Novel / held-out conditions: conditions the model never saw in training. We always test on unseen conditions.",
 "Baseline: a deliberately simple comparison method. If a complex model can't beat a simple baseline, the complexity isn't buying anything. Our key baseline is 'chemistry-similarity' (next slides).",
 "Gene-specific interaction: does the model learn something particular about how THIS gene reacts, versus echoing trends shared by all genes? That distinction is the crux.",
 "The 'ingredients' card: 1152-d embedding = each gene described by 1152 numbers, read from layer 8 ('L8'). 425-d multihot = a 425-slot on/off vector marking which chemicals are present (1 = present). relevance = max(0, −fit): fitness below zero means the knockout hurt growth, so the gene was needed → that condition is a 'stressor'; non-hurting conditions get relevance 0.",
],
6: [
 "The data: each row is one gene measured in one experiment (27.4 million such measurements), from the public RB-TnSeq 'Fitness Browser' dataset (Wetmore et al., 2015).",
 "Replicate structure: the same condition was run more than once, which lets us estimate measurement noise — later this becomes our performance 'ceiling'.",
 "Picture a matrix: rows = genes, columns = conditions, entries = fitness. CONDITION-HOLDOUT hides entire COLUMNS for testing. Those columns are 'cold' — zero observed entries.",
 "Matrix factorization / collaborative filtering: the 'Netflix recommendation' trick — fill missing matrix entries by learning hidden factors per row and column. Key limitation: it can only place a column that has SOME observed entries. A fully empty (cold) column can't be placed from the matrix alone.",
 "Side information / side-features: extra description of a column that lets you locate a never-seen one — here, the condition's CHEMISTRY. Because we must use it, this is 'inductive (cold-start) matrix completion', harder and more realistic than filling random missing cells.",
 "Stratified by stressor group: when splitting, we ensure each TYPE of stress (acids, metals, antibiotics, temperature…) appears on both sides, so the test set is representative.",
],
7: [
 "Spearman correlation: how well two orderings agree, −1 (reversed) to +1 (identical). Computed WITHIN each gene across its conditions, then averaged over genes.",
 "NDCG@5 (Normalized Discounted Cumulative Gain at 5): a ranking score from search engines. It rewards putting the truly biggest stressors in your top 5, with more credit for higher placement, normalized so 1.0 = a perfect top-5. We favor it because 'find the top stressors' is the real use-case.",
 "Bootstrap / confidence interval: resample the data many times to see how much the score wobbles. 'Hierarchical org→gene' resamples organisms first, then genes within them, because genes in the same organism aren't independent. 'CIs disjoint' = intervals don't overlap, so a difference is real, not luck.",
 "Promotion delta / 'the gate': before running anything we pre-registered the rule — a new method must beat the baseline by at least ~0.026 NDCG@5 to count as a win. That baseline is chem-kNN, 'the gate'.",
 "The ladder rows: chem-NULL = predict the average condition-effect across all genes (no gene specificity). deep model = our intended neural net. linear-MF = a gene vector learned purely from fitness data (a 'fitness-aware' representation). chem-kNN = look up the gene's own fitness at the most chemically-similar known conditions. ceiling = replicate agreement (the most any method could score).",
],
8: [
 "Three independent 'levers' you could pull to improve a model — all three fail to beat the lookup.",
 "Encoder: how we turn a condition's chemistry into numbers. 'Multihot' = which chemicals are present. 'Fingerprints' (Morgan, RDKit, MACCS) = standard cheminformatics encodings of a molecule's STRUCTURE as a bit-vector (which substructures it contains).",
 "Objective / loss function: the quantity minimized during training. MSE = mean squared error. Huber = like MSE but robust to outliers. RankNet / LambdaRank / ListMLE / ApproxNDCG = losses purpose-built to optimize ORDER rather than exact values.",
 "Capacity: how flexible/powerful the model is, roughly its number of learned 'parameters' (tunable numbers). A tiny linear model (~13k) ties a deep nonlinear net (~2.5M) — so raw power isn't the bottleneck.",
 "Global parametric model: one fixed set of weights applied to every gene (vs a 'local' lookup that answers per-gene). 'Plateau' = performance flattens no matter what we change — the tell-tale that the bottleneck is elsewhere.",
],
9: [
 "The WHY behind the result. Two kinds of signal: (1) the population / cross-gene pattern — trends shared by most genes (e.g., 'this antibiotic hurts almost everything'); and (2) the within-gene signal — the specific ordering of conditions for ONE gene.",
 "The population pattern alone carries essentially zero useful within-gene ordering (that's chem-NULL scoring near zero). So all the real signal is in each gene's OWN deviations from the crowd — its idiosyncrasies.",
 "A global model minimizes AVERAGE error over ~54,000 genes. To do that it learns the crowd behavior and treats each gene's quirks as random noise — it 'averages them away'. That's exactly the signal we needed.",
 "kNN (k-nearest neighbors): a 'non-parametric' method — no learned weights; to predict it finds the k most similar known examples and averages them. Here 'similar' = chemically-near conditions for the SAME gene, so it preserves that gene's idiosyncratic history.",
 "The doctor analogy: a global model is like diagnosing a stranger from their demographic 'type'; kNN is like having that exact patient's full chart. When the chart exists and the patient is unusual, the chart wins — not by being smarter, but by holding patient-specific information the global view discarded.",
],
10: [
 "If the model and the lookup each know different things, COMBINING them (a 'hybrid'/'ensemble') should beat either alone. We tried four ways; none clears the bar.",
 "z-scored: before blending two predictors we standardize each (subtract mean, divide by standard deviation) so they're on the same scale. Convex blend with weight α: final = α·kNN + (1−α)·model, α in [0,1]. 'Held-out α' = choose α on one half of the genes, report on the other half, so we don't cheat.",
 "Residual hybrid: train the model to predict only the part kNN gets WRONG (truth minus kNN), then add it back. 'LOO' (leave-one-out): when computing kNN's training-time prediction for a point, exclude that point itself so it can't copy its own answer.",
 "Retrieval-augmented: feed the model the kNN evidence (nearest neighbors' fitness) as extra inputs. Learned gating: train a small network to decide, case by case, how much to trust kNN vs the model.",
 "'Sign-flips by seed': rerunning with different random seeds gave sometimes +, sometimes − — the hallmark of noise, not a real effect. 'Promotable' = large AND reliable enough to count under our pre-registered rule. None were.",
],
11: [
 "This slide rebuts the most common objection: 'your answer key (the fitness labels) is noisy, so maybe the model is fine and the metric is just measuring noise'.",
 "Label noise: measurement error in the fitness values themselves — the ground truth we score against is imperfect.",
 "Moderated t-statistic (|t|): a per-measurement confidence score = effect size divided by its noise, then 'moderated' (stabilized by borrowing information across many measurements — the idea behind the limma method in genomics). Large |t| = a confident, real effect; |t| > 4 is the field's standard 'this is real' threshold. It's already in our dataset.",
 "Stratify into quartiles: sort genes by confidence and split into four equal groups (Q1 least … Q4 most confident), then analyze each separately.",
 "Reading the figure: as confidence rises (left→right) every method improves and so does the ceiling (0.59→0.83) — confirming noise really does depress scores. BUT the orange kNN line stays above the blue model line in every group; the gap shrinks (0.055→0.040) yet never closes. Even on the cleanest genes, the model loses — the deficit is structural, not just noise. Upweighting confident rows during training ('t-weighted') also didn't help (Δ ≈ −0.004).",
],
12: [
 "A consolidated recap. Two recurring contrasts to keep straight:",
 "Parametric vs non-parametric: the neural net learns fixed weights (parametric); the kNN lookup just stores and retrieves data (non-parametric). Here the simple non-parametric method wins.",
 "Fitness-blind vs fitness-aware representation: the frozen ProteomeLM embedding only knows sequence/structure — nothing about stress response ('fitness-blind'). The linear-MF gene vector is learned from fitness data ('fitness-aware') — yet it STILL loses to the lookup, telling us the bottleneck isn't the representation, it's the global-vs-local structure of the task.",
 "Memorization vs generalization: the task rewards recalling a known gene's own history (memorization) far more than learning a transferable rule (generalization) — which is why a lookup is hard to beat, and why cross-species transfer (the T-regime) failed.",
],
13: [
 "Characterization paper: a paper whose contribution is to rigorously map out and benchmark a problem — including a well-supported NEGATIVE result and a strong baseline — rather than to propose a new state-of-the-art method. These are valuable and publishable when done carefully.",
 "Pre-registered: we wrote success criteria and the analysis plan BEFORE running experiments, preventing 'moving the goalposts' after seeing results — which is what makes a negative result credible.",
 "Warm vs cold: 'warm' = the gene/condition was seen in training; 'cold' = never seen. Everything here is the warm-gene case. The cold-gene diagnostic tests genes never seen in training — and notably kNN has no history to look up for a brand-new gene, so it's the one regime where a global model could in principle help. We list it as optional, not a make-or-break gate.",
 "R-EMB / fitness-aware inductive embedding: the tempting 'moonshot' of training a new embedding that DOES encode fitness response. We defer it: linear-MF suggests it won't help the warm task, and its only real target — cold transfer across species — is exactly where the evidence is most discouraging (recall the T-regime).",
],
}

for _i, _sl in enumerate(prs.slides, start=1):
    if _i in NOTES:
        _ntf = _sl.notes_slide.notes_text_frame
        _ntf.text = NOTES[_i][0]
        for _para in NOTES[_i][1:]:
            _p = _ntf.add_paragraph(); _p.text = _para

prs.save("artifacts/deck/Conditional_Gene_Essentiality_lab.pptx")
print("saved:", len(prs.slides._sldIdLst), "slides, notes on", len(NOTES))
