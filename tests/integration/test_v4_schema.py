"""S0 first action: verify the v4 workbook has the expected schema.

If this fails, no later stage can run.
"""
from pathlib import Path

import openpyxl
import pytest


WORKBOOK = Path("data/media_composition_v4.xlsx")
SHEET = "Media_Components_ML"
EXPECTED_COLUMNS = [
    "Media",
    "Canonical_ID",
    "Compound_name",
    "Chemical_form",
    "Source_row_component",
    "Decomposition_type",
    "Ingredient_source",
    "Include_in_ml",
    "Source_dataset",
    "Source_url",
]


@pytest.mark.skipif(not WORKBOOK.exists(), reason="v4 workbook not present in this checkout")
def test_v4_has_media_components_ml_sheet():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    assert SHEET in wb.sheetnames, (
        f"v4 workbook missing required sheet '{SHEET}'. "
        f"Found: {wb.sheetnames}"
    )


@pytest.mark.skipif(not WORKBOOK.exists(), reason="v4 workbook not present in this checkout")
def test_v4_columns_match_contract():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    sh = wb[SHEET]
    header = next(sh.iter_rows(max_row=1, values_only=True))
    assert list(header) == EXPECTED_COLUMNS, (
        f"v4 column mismatch.\nExpected: {EXPECTED_COLUMNS}\nFound:    {list(header)}"
    )


@pytest.mark.skipif(not WORKBOOK.exists(), reason="v4 workbook not present in this checkout")
def test_v4_has_minimum_row_count():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    sh = wb[SHEET]
    # max_row includes header
    assert sh.max_row >= 4000, f"v4 row count too low: {sh.max_row}"
